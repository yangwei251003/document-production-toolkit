# Sanitized historical source; synthetic prose and portable workspace bindings.
from document_toolkit.runtime import enter_archive, workspace_root, source_path
enter_archive('career/assemble_v2.py')
from pathlib import Path
import json,shutil,hashlib,html,zipfile
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
from pypdf import PdfReader
R=workspace_root();V=R/'PPT制作_V2';D=R/'交付文件/求职综合展示PPT_V2_22页'
D.mkdir(exist_ok=True);(D/'逐页图片').mkdir(exist_ok=True)
O=json.loads((V/'outline.json').read_text('utf-8'));M=json.loads((V/'imagegen-manifest.json').read_text('utf-8'))
reviews={1:'合成示例内容：通过项目实践展示文档制作与检查流程/专业/合成示例内容：通过项目实践展示文档制作与检查流程',2:'合成示例内容：通过项目实践展示文档制作与检查流程',3:'合成示例内容：通过项目实践展示文档制作与检查流程',4:'合成示例内容：通过项目实践展示文档制作与检查流程',5:'合成示例内容：通过项目实践展示文档制作与检查流程',6:'合成示例内容：通过项目实践展示文档制作与检查流程',7:'合成示例内容：通过项目实践展示文档制作与检查流程',8:'公司/年月/合成示例内容：通过项目实践展示文档制作与检查流程',9:'两页整合，14/25/32合成示例内容：通过项目实践展示文档制作与检查流程',10:'马拉松2:56:31、羽毛球奖、45/160合成示例内容：通过项目实践展示文档制作与检查流程',12:'四步与个人/AI合成示例内容：通过项目实践展示文档制作与检查流程',14:'合成示例内容：通过项目实践展示文档制作与检查流程',15:'合成示例内容：通过项目实践展示文档制作与检查流程',16:'03合成示例内容：通过项目实践展示文档制作与检查流程',18:'在校至2028/初入岗位/合成示例内容：通过项目实践展示文档制作与检查流程',19:'合成示例内容：通过项目实践展示文档制作与检查流程/岗位完整，无联系方式。'}
for row in M:
 n=row['slide']
 if n in reviews: row.update(visual_review='PASS_WITH_SOURCE_LIMITS' if n==7 else 'PASS',review_notes=reviews[n])
 if n in [9,12] and row.get('prior_attempts'): row['prior_attempts'][-1]['visual_review']='REJECTED';row['prior_attempts'][-1]['review_notes']='出现无依据照片' if n==9 else '出现与项目无关的海洋装饰'
assert len(M)==22 and all(x['visual_review'].startswith('PASS') for x in M)
(V/'imagegen-manifest.json').write_text(json.dumps(M,ensure_ascii=False,indent=2),'utf-8')
for n,t in {9:'合成示例内容：通过项目实践展示文档制作与检查流程',12:'合成示例内容：通过项目实践展示文档制作与检查流程'}.items(): (V/'prompts'/f'{n:02}_revision.md').write_text(t,'utf-8')
p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
proof='合成示例内容：通过项目实践展示文档制作与检查流程'
script=['# 合成示例内容：通过项目实践展示文档制作与检查流程','','19合成示例内容：通过项目实践展示文档制作与检查流程','']
pdf=canvas.Canvas(str(D/'示例同学_求职综合展示_V2_全22页预览.pdf'),pagesize=(960,540))
for s in O['slides']:
 n=s['slide_number'];src=V/'slides'/f'{n:02}.png';dst=D/'逐页图片'/src.name;shutil.copy2(src,dst)
 slide=p.slides.add_slide(p.slide_layouts[6]);slide.shapes.add_picture(str(src),0,0,width=p.slide_width,height=p.slide_height)
 if s['hidden']:slide._element.set('show','0')
 note=s['speaker_draft']+'\n\n建议讲述秒数：'+str(s['speaker_seconds'])+'\n\n页面文字：\n'+'\n'.join(s['exact_text'])+'\n\n资料说明：'+proof
 if n==10:note+=' 合成示例内容：通过项目实践展示文档制作与检查流程'
 if n==5:note+=' 合成示例内容：通过项目实践展示文档制作与检查流程'
 slide.notes_slide.notes_text_frame.text=note
 pdf.drawImage(str(src),0,0,width=960,height=540);pdf.showPage()
 script+= [f'## {n:02} {s["title"]}'+('（答辩备用）' if s['hidden'] else ''),'',s['speaker_draft'],'','页面文字：','']+['- '+t for t in s['exact_text']]+['']
 s.update(generation_status='REUSED' if s['generation_status']=='REUSE' else 'GENERATED',review_status='VISUAL_REVIEWED')
pdf.save();ppt=D/'示例同学_求职综合展示_V2_22页.pptx';p.save(ppt)
assert ppt.stat().st_size<50_000_000
assert len(PdfReader(D/'示例同学_求职综合展示_V2_全22页预览.pdf').pages)==22
pres=Presentation(ppt);assert len(pres.slides)==22
assert all((s._element.get('show')=='0')==(i>19) for i,s in enumerate(pres.slides,1))
assert (V/'slides/17.png').read_bytes()==(R/'PPT制作/slides/16.png').read_bytes()
with zipfile.ZipFile(ppt) as z:
 media=[z.read(x) for x in z.namelist() if x.startswith('ppt/media/')];assert len(media)==22
 assert {hashlib.sha256(x).hexdigest() for x in media}=={r['sha256'] for r in M}
(D/'逐页文案与讲稿.md').write_text('\n'.join(script),'utf-8')
sources=(R/'交付文件/求职综合展示PPT/素材与资料来源.md').read_text('utf-8').split('## 原始图片')[0]
sources+=source_path('PowerGrid.html)\n\n上述为2026-10-23查询的行业与产品资料；企业对标为求职研究判断，不意味着当前岗位开放或用户已符合届别。云圣机巢素材是官方产品视觉，不能称为现场实拍。\n\n')
sources+='合成示例内容：通过项目实践展示文档制作与检查流程\n\n## V2逐页图片来源\n\n'
for s in O['slides']:
 if s['reference_paths']: sources+=f"### 第{s['slide_number']}页\n\n"+'\n'.join('- '+x for x in s['reference_paths'])+'\n\n'
(D/'素材与资料来源.md').write_text(sources,'utf-8')
for name in ['imagegen-manifest.json','修订说明.md']:shutil.copy2(V/name,D/('图片生成与审查记录.json' if name.endswith('json') else name))
readme=f'# 合成示例内容：通过项目实践展示文档制作与检查流程\n\n- 合成示例内容：通过项目实践展示文档制作与检查流程\n- PPTX大小：{ppt.stat().st_size / 1000000.0:.2f} MB，合成示例内容：通过项目实践展示文档制作与检查流程\n- 合成示例内容：通过项目实践展示文档制作与检查流程\n- 合成示例内容：通过项目实践展示文档制作与检查流程\n- 合成示例内容：通过项目实践展示文档制作与检查流程\n- 合成示例内容：通过项目实践展示文档制作与检查流程\n- 合成示例内容：通过项目实践展示文档制作与检查流程\n- 合成示例内容：通过项目实践展示文档制作与检查流程\n'
(D/'使用说明.md').write_text(readme,'utf-8')
cards=''.join(f'<section><h2>{s["slide_number"]:02} {html.escape(s["title"])}'+(' · 答辩备用' if s['hidden'] else '')+f'</h2><img loading="lazy" src="逐页图片/{s["slide_number"]:02}.png"><p>{html.escape(s["speaker_draft"])}</p></section>' for s in O['slides'])
(D/'全套浏览.html').write_text('<!doctype html><meta charset="utf-8"><title>合成示例内容：通过项目实践展示文档制作与检查流程</title><style>body{margin:30px auto;max-width:1280px;background:#10283f;color:#f7f8fa;font-family:Microsoft YaHei,sans-serif}h1{font-size:30px}section{margin:40px 0}h2{font-size:20px}img{width:100%;box-shadow:0 8px 24px #0005}p{line-height:1.8;color:#c8d4df}</style><h1>合成示例内容：通过项目实践展示文档制作与检查流程</h1><p>19合成示例内容：通过项目实践展示文档制作与检查流程</p>'+cards,'utf-8')
O['stage']='ASSEMBLED_PENDING_NATIVE_QA';(V/'outline.json').write_text(json.dumps(O,ensure_ascii=False,indent=2),'utf-8')
report={'status':'STRUCTURAL_PASS_PENDING_NATIVE','slides':22,'main_slides':19,'hidden_backup_slides':3,'pptx_bytes':ppt.stat().st_size,'pptx_MB_decimal':round(ppt.stat().st_size/1e6,2),'pdf_pages':22,'old16_equals_new17':True,'image_hashes_match_pptx':True,'planned_seconds':335,'timed_rehearsal':False}
(V/'qa/交付检查.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
ps=(R/'工作记录/verify_ppt_native.ps1').read_text('utf-8').replace('交付文件\\求职综合展示PPT\\示例同学_求职综合展示_电力能源方向.pptx','交付文件\\求职综合展示PPT_V2_22页\\示例同学_求职综合展示_V2_22页.pptx').replace('PPT制作\\qa','PPT制作_V2\\qa');(R/'工作记录/verify_v2_native.ps1').write_text(ps,'utf-8-sig')
print(json.dumps(report,ensure_ascii=False))
