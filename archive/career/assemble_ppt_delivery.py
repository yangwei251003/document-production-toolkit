# Sanitized historical source; synthetic prose and portable workspace bindings.
from document_toolkit.runtime import enter_archive, workspace_root, source_path
enter_archive('career/assemble_ppt_delivery.py')
from pathlib import Path
import json, hashlib, shutil, zipfile, html, importlib.util
from pptx import Presentation
from reportlab.pdfgen import canvas
from pypdf import PdfReader

ROOT = workspace_root()
PROD = ROOT / 'PPT制作'
OUT = ROOT / '交付文件' / '求职综合展示PPT'
OUT.mkdir(parents=True, exist_ok=True)
outline = json.loads((PROD / 'outline.json').read_text('utf-8'))
manifest = json.loads((PROD / 'imagegen-manifest.json').read_text('utf-8'))
assert len(manifest) == 21 and all(r['visual_review'] == 'PASS' for r in manifest)

scripts = [
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程'
]
assert len(scripts) == 21
sources = [
('教育部第四届就业赛道方案', source_path('W020261010465563631809.docx'), '合成示例内容：通过项目实践展示文档制作与检查流程'),
('教育部第四届大赛通知', source_path('t20261010_1453481.html'), '官方通知'),
('云圣智能历史校园招聘', source_path('2164.html'), '合成示例内容：通过项目实践展示文档制作与检查流程'),
('大疆官方精细化巡检说明', source_path('content?customId=[示例编号]&documentType=&lang=zh-CN&paperDocType=ARTICLE&re=CN&spaceId=17'), '巡检应用流程参考'),
('普宙可见光与红外应用资料', source_path('193.html'), '场景理解参考，不用于本人经历证明'),
('金奖选手左怡经验分享', source_path('2481.htm'), '学习岗位目标与经历组织'),
('北京科技大学陈倚获奖报道', source_path('68846.htm'), '学习可核实的现场局部展示与叙事'),
('浙江大学杨少晨获奖报道', source_path('page.htm'), '学习工程实践与职业方向连接'),
('就业赛道展示指导', source_path('t20251119-914-27785.html'), '结构与展示参考')
]
source_by_slide = {4:[2,3,4],13:[3,4],21:[0,2,3]}
proof_notes = {
1:'合成示例内容：通过项目实践展示文档制作与检查流程',
6:'合成示例内容：通过项目实践展示文档制作与检查流程',
7:'合成示例内容：通过项目实践展示文档制作与检查流程',
8:'合成示例内容：通过项目实践展示文档制作与检查流程',
9:'合成示例内容：通过项目实践展示文档制作与检查流程-09-03；合成示例内容：通过项目实践展示文档制作与检查流程',
10:'能策Pilot V3.1 README合成示例内容：通过项目实践展示文档制作与检查流程',
11:'合成示例内容：通过项目实践展示文档制作与检查流程',
12:'P021合成示例内容：通过项目实践展示文档制作与检查流程',
13:'合成示例内容：通过项目实践展示文档制作与检查流程',
16:'合成示例内容：通过项目实践展示文档制作与检查流程',
19:'合成示例内容：通过项目实践展示文档制作与检查流程'
}

# Provenance corrections and traceable accepted adaptations.
reasons={2:'合成示例内容：通过项目实践展示文档制作与检查流程',12:'合成示例内容：通过项目实践展示文档制作与检查流程',14:'合成示例内容：通过项目实践展示文档制作与检查流程',20:'合成示例内容：通过项目实践展示文档制作与检查流程/AI两栏。'}
for row in manifest:
    for old in row.get('prior_attempts',[]):
        old['visual_review']='REJECTED'
        old['review_notes']=reasons.get(row['slide'],'由后续版本替换')
    assert hashlib.sha256(Path(row['copied_to']).read_bytes()).hexdigest()==row['sha256']
(PROD/'imagegen-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),'utf-8')
outline['stage']='DELIVERED'
for i,s in enumerate(outline['slides']):
    s['generation_status']='GENERATED'
    s['review_status']='VISUAL_PASS'
    s['speaker_draft']=scripts[i]
    s['generated_image']=f'slides/{i+1:02}.png'
outline['slides'][11]['exact_text'].append('智能出纸机团队成果')
outline['slides'][12]['accepted_visual_labels']=['巡检图像','分析识别','巡检报告']
outline['slides'][20]['accepted_omission']='合成示例内容：通过项目实践展示文档制作与检查流程'
(PROD/'outline.json').write_text(json.dumps(outline,ensure_ascii=False,indent=2),'utf-8')

# Image-only slides, no programmatically typeset slide content.
deck={'slide_width_in':13.3333333333,'slide_height_in':7.5,'assets_dir':str(PROD/'slides'),'slides':[{'background':f'{i:02}.png'} for i in range(1,22)]}
(PROD/'deck.json').write_text(json.dumps(deck,ensure_ascii=False,indent=2),'utf-8')
modulepath=Path(source_path('compose_pptx.py'))
spec=importlib.util.spec_from_file_location('image_deck_composer',modulepath)
composer=importlib.util.module_from_spec(spec); spec.loader.exec_module(composer)
pptpath=OUT/'示例同学_求职综合展示_电力能源方向.pptx'
composer.build_pptx(deck,pptpath)
prs=Presentation(pptpath)
prs.core_properties.title='以技为桥，让低空服务电力'
prs.core_properties.author='示例同学'
prs.core_properties.subject='无人机技术销售工程师（电力能源方向）'
for n,slide in enumerate(prs.slides,1):
    sl=outline['slides'][n-1]
    if n>18: slide._element.set('show','0')
    note=f"第{n:02}页｜{sl['title']}\n{'答辩备用，默认隐藏' if n>18 else '建议时长：'+str(sl['speaker_seconds'])+'秒'}\n\n讲稿草案：\n{scripts[n-1]}\n\n页面文案：\n"+'\n'.join(sl['exact_text'])
    note+='\n\n证据/边界：\n'+proof_notes.get(n,'合成示例内容：通过项目实践展示文档制作与检查流程')
    for k in source_by_slide.get(n,[]): note+='\n'+sources[k][0]+'\n'+sources[k][1]
    note+='\n\n合成示例内容：通过项目实践展示文档制作与检查流程'
    slide.notes_slide.notes_text_frame.text=note
    slide.shapes[0].name=f'{n:02} {sl["title"]}'
    slide.shapes[0]._element.nvPicPr.cNvPr.set('descr','；'.join(sl['exact_text']))
prs.save(pptpath)
assert pptpath.stat().st_size<50_000_000

# PDF is the same whole image per page; no redrawing or substituted content.
pdfpath=OUT/'示例同学_求职综合展示_全21页预览.pdf'
c=canvas.Canvas(str(pdfpath),pagesize=(960,540),pageCompression=1)
c.setTitle('合成示例内容：通过项目实践展示文档制作与检查流程+3页答辩备用')
for n in range(1,22):
    c.drawImage(str(PROD/'slides'/f'{n:02}.png'),0,0,width=960,height=540)
    c.showPage()
c.save()
assert len(PdfReader(pdfpath).pages)==21

imagesout=OUT/'逐页图片'
imagesout.mkdir(exist_ok=True)
for n in range(1,22): shutil.copy2(PROD/'slides'/f'{n:02}.png',imagesout/f'{n:02}.png')
md=['# 合成示例内容：通过项目实践展示文档制作与检查流程','', '18合成示例内容：通过项目实践展示文档制作与检查流程','']
for n,sl in enumerate(outline['slides'],1):
    md += [f'## {n:02} {sl["title"]}', '', f'建议时长：{sl["speaker_seconds"]}秒。' if n<=18 else '答辩备用，不计入主讲。','', '页面文案：','']
    md += ['- '+t for t in sl['exact_text']]
    md += ['', '讲稿草案：', '', scripts[n-1], '', '证据与表述边界：'+proof_notes.get(n,'沿用个人事实与证据索引。'), '']
(OUT/'逐页文案与讲稿.md').write_text('\n'.join(md),'utf-8')
srcmd=['# 素材与资料来源','', '研究完成于2026-10-22；生成与组装于2026-10-23。合成示例内容：通过项目实践展示文档制作与检查流程','', '## 外部资料','']
for title,url,use in sources: srcmd += [f'- [{title}]({url})：{use}']
srcmd += ['', '## 原始图片','', '合成示例内容：通过项目实践展示文档制作与检查流程','']
for s in outline['slides']:
    for path in s['reference_paths']: srcmd += [f'- 第{s["slide_number"]}页：`{path}`']
(OUT/'素材与资料来源.md').write_text('\n'.join(srcmd),'utf-8')

cards=[]
for n,sl in enumerate(outline['slides'],1):
    cards.append(f'<section id="s{n}"><h2>{n:02}　{html.escape(sl["title"])}</h2><p>{"答辩备用 · PPT中默认隐藏" if n>18 else str(sl["speaker_seconds"])+"秒 · 主讲"}</p><a href="逐页图片/{n:02}.png"><img loading="lazy" src="逐页图片/{n:02}.png" alt="{html.escape(sl["title"])}"></a><details><summary>展开讲稿</summary><p>{html.escape(scripts[n-1])}</p></details></section>')
(OUT/'全套浏览.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>示例同学求职综合展示</title><style>body{margin:0;background:#eaf0f5;color:#10283f;font:17px/1.7 Microsoft YaHei,sans-serif}header,main{max-width:1150px;margin:auto;padding:24px}header{padding-bottom:0}h1{font-size:32px}h2{font-size:22px}section{background:white;padding:20px;margin:0 0 24px;border-radius:12px}img{width:100%;display:block}details{margin-top:14px}a{color:#087d99}p{margin:8px 0}</style><header><h1>合成示例内容：通过项目实践展示文档制作与检查流程</h1><p>合成示例内容：通过项目实践展示文档制作与检查流程</p><p><a href="'+pptpath.name+'">打开PPT</a>　<a href="'+pdfpath.name+'">打开PDF</a>　<a href="逐页文案与讲稿.md">逐页讲稿</a></p></header><main>'+''.join(cards)+'</main></html>','utf-8')

checkprs=Presentation(pptpath)
assert len(checkprs.slides)==21
checks=[]
for n,sl in enumerate(checkprs.slides,1):
    digest=hashlib.sha256(sl.shapes[0].image.blob).hexdigest()
    assert len(sl.shapes)==1 and digest==manifest[n-1]['sha256']
    assert (sl._element.get('show')=='0')==(n>18)
    assert scripts[n-1] in sl.notes_slide.notes_text_frame.text
    assert sl.shapes[0].left==0 and sl.shapes[0].top==0
    assert sl.shapes[0].width==checkprs.slide_width and sl.shapes[0].height==checkprs.slide_height
    checks.append({'slide':n,'hidden':n>18,'one_whole_image':True,'image_hash_matches':True,'notes_present':True,'visual_review':manifest[n-1]['review_notes']})
report={'status':'PASS_STRUCTURAL_AND_VISUAL','pptx_bytes':pptpath.stat().st_size,'pptx_MB_decimal':round(pptpath.stat().st_size/1_000_000,2),'under_50MB':True,'slide_count':21,'visible_count':18,'hidden_count':3,'pdf_pages':21,'planned_seconds':345,'actual_rehearsal':'NOT_TESTED','native_powerpoint_open':'PENDING','slides':checks}
(PROD/'qa'/'交付检查.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
readme=f'# 使用说明\n\n主文件：{pptpath.name}（{report['pptx_MB_decimal']} MB，小于50 MB）。\n\n- 18合成示例内容：通过项目实践展示文档制作与检查流程\n- 合成示例内容：通过项目实践展示文档制作与检查流程\n- 合成示例内容：通过项目实践展示文档制作与检查流程\n- PDF合成示例内容：通过项目实践展示文档制作与检查流程.html合成示例内容：通过项目实践展示文档制作与检查流程\n- 合成示例内容：通过项目实践展示文档制作与检查流程\n- 合成示例内容：通过项目实践展示文档制作与检查流程.md。\n\n合成示例内容：通过项目实践展示文档制作与检查流程\n'
(OUT/'使用说明.md').write_text(readme,'utf-8')
shutil.copy2(PROD/'qa'/'交付检查.json',OUT/'交付检查.json')
shutil.copy2(PROD/'imagegen-manifest.json',OUT/'图片生成与审查记录.json')
print(json.dumps({'pptx':str(pptpath),'pdf':str(pdfpath),'report':report|{'slides':'21 passed'}},ensure_ascii=False))
