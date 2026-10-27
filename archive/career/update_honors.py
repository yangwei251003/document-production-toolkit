# Sanitized historical source; synthetic prose and portable workspace bindings.
from document_toolkit.runtime import enter_archive, workspace_root, source_path
enter_archive('career/update_honors.py')
from pathlib import Path
import json,shutil,hashlib,zipfile
from pptx import Presentation
from reportlab.pdfgen import canvas
from pypdf import PdfReader
from PIL import Image
R=workspace_root();V=R/'PPT制作_V2';D=R/'交付文件/求职综合展示PPT_V2_22页';B=V/'版本备份/荣誉页扩充前'
B.mkdir(parents=True,exist_ok=True)
C=json.loads((V/'荣誉页扩充配置.json').read_text('utf-8'));ppt=D/'示例同学_求职综合展示_V2_22页.pptx'
for src in [ppt,V/'slides/07.png',V/'outline.json',V/'imagegen-manifest.json']:
 dest=B/src.name
 if not dest.exists():shutil.copy2(src,dest)
oldhash=hashlib.sha256((V/'slides/07.png').read_bytes()).hexdigest()
shutil.copy2(C['source'],V/'slides/07.png');shutil.copy2(C['source'],D/'逐页图片/07.png')
shutil.copy2(C['sheet_source'],V/'assets/新增四证书拼版_imagegen.png')
(V/'prompts/07_eight_certificates_final.md').write_text(C['prompt'],'utf-8')
O=json.loads((V/'outline.json').read_text('utf-8'));s=O['slides'][6]
s['reference_paths']=C['references'];s['prompt_file']=str(V/'prompts/07_eight_certificates_final.md')
s['design_spec_markdown']='合成示例内容：通过项目实践展示文档制作与检查流程'
s['exact_text']=['以持续学习，积累专业底气','2025—2026学年 国家奖学金','2026 合成示例内容：通过项目实践展示文档制作与检查流程','2025 合成示例内容：通过项目实践展示文档制作与检查流程','2026 合成示例内容：通过项目实践展示文档制作与检查流程','2026 合成示例内容：通过项目实践展示文档制作与检查流程','2025 合成示例内容：通过项目实践展示文档制作与检查流程','2024—2025学年 学校一等奖学金','2024—2025学年 优秀学生干部','2024—2025学年 优秀五好学生','合成示例内容：通过项目实践展示文档制作与检查流程','合成示例内容：通过项目实践展示文档制作与检查流程']
s['speaker_draft']='合成示例内容：通过项目实践展示文档制作与检查流程'
s['revision']='2026-10-23 荣誉页3张扩充为8张'
(V/'outline.json').write_text(json.dumps(O,ensure_ascii=False,indent=2),'utf-8')
M=json.loads((V/'imagegen-manifest.json').read_text('utf-8'));row=next(x for x in M if x['slide']==7);prev=dict(row)
row.update(generated_source=C['source'],sha256=hashlib.sha256((V/'slides/07.png').read_bytes()).hexdigest(),dimensions=list(Image.open(V/'slides/07.png').size),prompt_file=s['prompt_file'],reference_paths=C['references'],intermediate_generated_source=C['sheet_source'],visual_review='PASS_WITH_SOURCE_LIMITS',review_notes='合成示例内容：通过项目实践展示文档制作与检查流程/合成示例内容：通过项目实践展示文档制作与检查流程',prior_attempts=row.get('prior_attempts',[])+[prev])
(V/'imagegen-manifest.json').write_text(json.dumps(M,ensure_ascii=False,indent=2),'utf-8');shutil.copy2(V/'imagegen-manifest.json',D/'图片生成与审查记录.json')
p=Presentation(ppt);slide=p.slides[6];shape=slide.shapes[0];rid=shape._element.blipFill.blip.rEmbed;part=slide.part.related_part(rid);part._blob=(V/'slides/07.png').read_bytes()
slide.notes_slide.notes_text_frame.text=s['speaker_draft']+'\n\n页面文字：\n'+'\n'.join(s['exact_text'])+'\n\n合成示例内容：通过项目实践展示文档制作与检查流程'
p.save(ppt);assert ppt.stat().st_size<50_000_000
with zipfile.ZipFile(B/ppt.name) as a,zipfile.ZipFile(ppt) as b:
 oldimgs={hashlib.sha256(a.read(n)).hexdigest() for n in a.namelist() if n.startswith('ppt/media/')};newimgs={hashlib.sha256(b.read(n)).hexdigest() for n in b.namelist() if n.startswith('ppt/media/')}
 assert len(oldimgs & newimgs)==21 and len(newimgs)==22
pdf=canvas.Canvas(str(D/'示例同学_求职综合展示_V2_全22页预览.pdf'),pagesize=(960,540))
for n in range(1,23):pdf.drawImage(str(V/'slides'/f'{n:02}.png'),0,0,width=960,height=540);pdf.showPage()
pdf.save();assert len(PdfReader(D/'示例同学_求职综合展示_V2_全22页预览.pdf').pages)==22
md=D/'逐页文案与讲稿.md';text=md.read_text('utf-8');start=text.index('## 07 ');end=text.index('## 08 ',start)
text=text[:start]+'## 07 '+s['title']+'\n\n'+s['speaker_draft']+'\n\n页面文字：\n\n'+'\n'.join('- '+x for x in s['exact_text'])+'\n\n'+text[end:];md.write_text(text,'utf-8')
h=D/'全套浏览.html';t=h.read_text('utf-8');t=t.replace('逐页图片/07.png','逐页图片/07.png?v=8cert');start=t.index('<section><h2>07 ');end=t.index('</section>',start)
t=t[:start]+f'<section><h2>07 {s["title"]} · 八证书扩充版</h2><img loading="lazy" src="逐页图片/07.png?v=8cert"><p>{s["speaker_draft"]}</p>'+t[end:];h.write_text(t,'utf-8')
for f in ['使用说明.md','逐页检查与使用说明.md']:
 path=D/f;text=path.read_text('utf-8').replace('40.36',f'{ppt.stat().st_size/1e6:.2f}').replace('三张证书参考图及三项资格','八张证书参考图及三项资格');text+='\n\n2026-10-23合成示例内容：通过项目实践展示文档制作与检查流程\n';path.write_text(text,'utf-8')
with (D/'素材与资料来源.md').open('a',encoding='utf-8') as f:f.write('\n## 第7页八证书扩充原件\n\n'+'\n'.join('- '+x for x in C['references'])+'\n\n合成示例内容：通过项目实践展示文档制作与检查流程\n')
report=json.loads((V/'qa/交付检查.json').read_text('utf-8'));report.update(status='HONOR_REVISION_PENDING_NATIVE',pptx_bytes=ppt.stat().st_size,pptx_MB_decimal=round(ppt.stat().st_size/1e6,2),honor_certificates=8,other_21_images_unchanged=True)
(V/'qa/交付检查.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
print(json.dumps({'pptx_MB':report['pptx_MB_decimal'],'certificates':8,'other_21_images_unchanged':True},ensure_ascii=False))
