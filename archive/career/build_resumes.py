# Sanitized historical source; synthetic prose and portable workspace bindings.
from document_toolkit.runtime import enter_archive, workspace_root, source_path
enter_archive('career/build_resumes.py')
from pathlib import Path
import json, sys
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

sys.stdout.reconfigure(encoding='utf-8')
ROOT = workspace_root()
OUT = ROOT / '交付文件'
PHOTO = Path(source_path('fixture-portrait.png'))
FONT = 'Microsoft YaHei'
BLUE = '274B60'
BODY = '20262B'

def font(run, size=10.5, bold=False, color=BODY):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    pr = run._element.get_or_add_rPr()
    rf = pr.rFonts
    for k in ['ascii','hAnsi','eastAsia','cs']:
        rf.set(qn('w:'+k), FONT)

def rich(p, text, size=10.5, color=BODY):
    for i,part in enumerate(text.split('**')):
        font(p.add_run(part),size,i%2==1,color)

def para(doc,text='',size=10.5,after=2,before=0,bold=False,color=BODY,keep=False):
    p=doc.add_paragraph()
    f=p.paragraph_format
    f.space_before=Pt(before); f.space_after=Pt(after)
    f.line_spacing=Pt(15.5 if size<=11 else size+5)
    f.keep_with_next=keep; f.widow_control=True
    rich(p,text if not bold else '**'+text+'**',size,color)
    return p

def bullet(doc,text,after=2):
    p=para(doc,text,after=after)
    pr=p._p.get_or_add_pPr()
    num=OxmlElement('w:numPr')
    il=OxmlElement('w:ilvl'); il.set(qn('w:val'),'0')
    ni=OxmlElement('w:numId'); ni.set(qn('w:val'),'88')
    num.append(il);num.append(ni);pr.append(num)
    p.paragraph_format.left_indent=Mm(3.5)
    p.paragraph_format.first_line_indent=Mm(-3.5)
    return p

def section(doc,title):
    p=para(doc,title,size=12,after=4,before=8,bold=True,color='000000',keep=True)
    p.style=doc.styles['Heading 1']
    return p

def entry(doc,title,date=''):
    p=para(doc,after=2,keep=True)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(178),WD_TAB_ALIGNMENT.RIGHT)
    font(p.add_run(title),10.7,True,'000000')
    if date: font(p.add_run('\t'+date),9.5,False,'58636C')
    return p

def picture_anchor(p):
    inline=p.add_run().add_picture(str(PHOTO),width=Mm(25))._inline
    inline.docPr.set('descr','示例同学本人正装证件照')
    anchor=OxmlElement('wp:anchor')
    for k,v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'1','behindDoc':'0','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items(): anchor.set(k,v)
    s=OxmlElement('wp:simplePos');s.set('x','0');s.set('y','0');anchor.append(s)
    h=OxmlElement('wp:positionH');h.set('relativeFrom','margin');o=OxmlElement('wp:posOffset');o.text=str(int(Mm(153)));h.append(o);anchor.append(h)
    v=OxmlElement('wp:positionV');v.set('relativeFrom','paragraph');o=OxmlElement('wp:posOffset');o.text='0';v.append(o);anchor.append(v)
    for name in ['extent','effectExtent']:
        el=inline.find(qn('wp:'+name))
        if el is not None: anchor.append(el)
    anchor.append(OxmlElement('wp:wrapNone'))
    for name in ['docPr','cNvGraphicFramePr']:
        el=inline.find(qn('wp:'+name))
        if el is not None: anchor.append(el)
    anchor.append(inline.find(qn('a:graphic')))
    inline.getparent().replace(inline,anchor)

def new_doc(target,detail=False):
    d=Document();s=d.sections[0]
    s.page_width=Mm(210);s.page_height=Mm(297)
    s.left_margin=Mm(16);s.right_margin=Mm(16);s.top_margin=Mm(14);s.bottom_margin=Mm(13)
    s.header_distance=Mm(5);s.footer_distance=Mm(6)
    for style in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
        st=d.styles[style];st.font.name=FONT;st.font.size=Pt(10.5);st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().append(parse_xml('<w:rFonts '+nsdecls('w')+' w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei" w:eastAsia="Microsoft YaHei"/>'))
        st.paragraph_format.space_after=Pt(2)
    normal=d.styles['Normal'];normal.paragraph_format.line_spacing=Pt(15.5)
    numbering=d.part.numbering_part.element
    abstract=parse_xml('<w:abstractNum '+nsdecls('w')+' w:abstractNumId="88"><w:multiLevelType w:val="singleLevel"/><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/><w:lvlJc w:val="left"/><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/></w:rPr></w:lvl></w:abstractNum>')
    numbering.append(abstract);numbering.append(parse_xml('<w:num '+nsdecls('w')+' w:numId="88"><w:abstractNumId w:val="88"/></w:num>'))
    p=para(d,'示例同学',size=26,after=4,bold=True,color='000000');p.style=d.styles['Title'];picture_anchor(p)
    p.paragraph_format.line_spacing=Pt(31)
    para(d,'求职意向  '+target,size=11,after=3,bold=True,color='000000')
    para(d,'[示例电话]   ·   student@example.com',size=10,after=4,color='46535D')
    para(d,'国家奖学金获得者  |  2025—2026学年',size=11,after=7,bold=True,color=BLUE)
    if detail:
        foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        font(foot.add_run('示例同学  ·  经历详细版    '),8.5,color='737D84')
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    d.core_properties.author='示例同学';d.core_properties.last_modified_by='示例同学'
    d.core_properties.subject='求职简历';d.core_properties.keywords='无人机 技术销售 电力能源 AI协作'
    return d

def education(d):
    section(d,'教育背景')
    entry(d,'示例大学  |  本科','2024.09—2028（预计）')
    para(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
    para(d,'2025—2026学年：**合成示例内容：通过项目实践展示文档制作与检查流程**（同年级同专业）。')

def internship(d,variant='main',detail=False):
    section(d,'实习经历')
    entry(d,'示例科技单位','2026.09—至今')
    para(d,'实习生  |  合成示例内容：通过项目实践展示文档制作与检查流程',size=10,color='58636C')
    if variant=='overseas':
        items=['合成示例内容：通过项目实践展示文档制作与检查流程','合成示例内容：通过项目实践展示文档制作与检查流程','合成示例内容：通过项目实践展示文档制作与检查流程']
    elif variant=='power':
        items=['合成示例内容：通过项目实践展示文档制作与检查流程','合成示例内容：通过项目实践展示文档制作与检查流程','合成示例内容：通过项目实践展示文档制作与检查流程']
    else:
        items=['合成示例内容：通过项目实践展示文档制作与检查流程','合成示例内容：通过项目实践展示文档制作与检查流程']
    for x in items:bullet(d,x)
    if detail:
        bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')

ENERGY=[
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程']
ROUTER=[
'合成示例内容：通过项目实践展示文档制作与检查流程',
'合成示例内容：通过项目实践展示文档制作与检查流程']

def projects(d,variant='main',detail=False):
    section(d,'项目经历')
    if variant=='overseas':
        entry(d,'云游地图  |  文旅网站','2026')
        bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
        bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
        entry(d,'动态技能路由器  |  AI 工作流与插件','2026')
        for x in ROUTER:bullet(d,x)
        return
    entry(d,'能策 Pilot  |  光储方案工作台','2026')
    for x in ENERGY:bullet(d,x)
    if variant=='power':
        entry(d,'智能出纸机  |  智能家电创新项目','2026')
        bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
        bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
        return
    if detail:bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
    entry(d,'动态技能路由器  |  AI 工作流与插件','2026')
    for x in ROUTER:bullet(d,x)
    if detail:bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')

def leadership(d,detail=False):
    section(d,'组织与实践')
    if detail:
        entry(d,'GKD视界之翼无人机协会  |  会长','2026.09—至今')
        bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
        entry(d,'三下乡社会实践队  |  队长','2026.07')
        bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
        bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
        entry(d,'24电气工程本科5班  |  班长','在校期间')
        bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
    else:
        bullet(d,'**无人机协会会长**：合成示例内容：通过项目实践展示文档制作与检查流程')
        bullet(d,'**三下乡队长**：合成示例内容：通过项目实践展示文档制作与检查流程')

def awards(d,detail=False,overseas=False):
    section(d,'专业竞赛' if detail else '代表荣誉')
    items=[
('2026','合成示例内容：通过项目实践展示文档制作与检查流程**三等奖**'),
('2026','合成示例内容：通过项目实践展示文档制作与检查流程**二等奖**'),
('2025','合成示例内容：通过项目实践展示文档制作与检查流程**全国总决赛一等奖**')]
    for date,t in items:para(d,date+'  '+t)
    if detail:
        para(d,'2025  合成示例内容：通过项目实践展示文档制作与检查流程**合成示例内容：通过项目实践展示文档制作与检查流程**')
        section(d,'文体与综合荣誉')
        para(d,'2026  《合成示例内容：通过项目实践展示文档制作与检查流程**二等奖**')
        para(d,'2025  广州马拉松净计时 **2小时56分31秒**；源羽杯青年组男子双打 **第一名**')
        para(d,'2024—2025合成示例内容：通过项目实践展示文档制作与检查流程')

def skills(d,overseas=False,detail=False):
    section(d,'技能与资格')
    bullet(d,'**AI 应用**：合成示例内容：通过项目实践展示文档制作与检查流程')
    bullet(d,'**资格证书**：合成示例内容：通过项目实践展示文档制作与检查流程')
    if overseas:para(d,'**海外意愿**：合成示例内容：通过项目实践展示文档制作与检查流程')
    if detail:para(d,'**应用方向**：合成示例内容：通过项目实践展示文档制作与检查流程')

def compact(kind):
    targets={'main':'无人机技术销售工程师（电力能源方向）','power':'无人机技术销售 / 合成示例内容：通过项目实践展示文档制作与检查流程','overseas':'合成示例内容：通过项目实践展示文档制作与检查流程/ 技术支持'}
    d=new_doc(targets[kind]);education(d);internship(d,kind);projects(d,kind)
    leadership(d);awards(d);skills(d,kind=='overseas')
    return d

def detailed():
    d=new_doc('无人机技术销售工程师（电力能源方向）',True)
    education(d);internship(d,detail=True);projects(d,detail=True)
    section(d,'相关项目')
    entry(d,'云游地图  |  文旅网站','2026')
    bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
    bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
    d.add_page_break()
    para(d,'示例同学',size=17,bold=True,color='000000',after=5)
    para(d,'工程实践与综合表现',size=10,color='58636C',after=3)
    section(d,'工程实践')
    entry(d,'智能出纸机  |  智能家电创新项目','2026')
    bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
    bullet(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
    leadership(d,True);awards(d,True);skills(d,detail=True)
    section(d,'职业准备')
    para(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
    para(d,'合成示例内容：通过项目实践展示文档制作与检查流程')
    return d

names={'main':'01_示例同学_比赛商务版_示例','detail':'02_示例同学_经历详细版_示例','power':'03_示例同学_电力企业定向版_示例','overseas':'04_示例同学_海外方向定向版_示例'}
OUT.mkdir(exist_ok=True)
for key,stem in names.items():
    d=detailed() if key=='detail' else compact(key)
    d.core_properties.title=stem.replace('_',' ')
    for element in list(d.styles.element.iter(qn('w:pBdr'))) + list(d.element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    d.save(OUT/(stem+'.docx'))
    (ROOT/'工作记录'/(stem+'.txt')).write_text('\n'.join(p.text for p in d.paragraphs),encoding='utf-8')
    print(stem,len(''.join(p.text for p in d.paragraphs)))
(ROOT/'工作记录'/'文件版本.json').write_text(json.dumps(names,ensure_ascii=False,indent=2),encoding='utf-8')
