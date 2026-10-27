# Sanitized historical source; synthetic prose and portable workspace bindings.
from document_toolkit.runtime import enter_archive, workspace_root, source_path
enter_archive('career/render_and_check.py')
from pathlib import Path
import sys, os, importlib.util, json, hashlib, re
import pdfplumber
from docx import Document
sys.stdout.reconfigure(encoding='utf-8')
ROOT=workspace_root()
RT=Path(source_path('codex-primary-runtime'))
SCRIPT=Path(source_path('render_docx.py'))
os.environ['PATH']=str(RT/'dependencies/native/poppler/Library/bin')+os.pathsep+os.environ.get('PATH','')
spec=importlib.util.spec_from_file_location('canonical_renderer',SCRIPT)
r=importlib.util.module_from_spec(spec);spec.loader.exec_module(r)
def native_word_pdf(doc_path,*args,**kwargs):
    # Word is the conversion backend; only freshly exported paired PDFs are accepted.
    p=Path(doc_path).with_suffix('.pdf')
    if not p.exists() or p.stat().st_mtime < Path(doc_path).stat().st_mtime:
        raise RuntimeError('Native Word PDF missing or older than source DOCX')
    return str(p),'Native Microsoft Word ExportAsFixedFormat; no LibreOffice invoked.'
r.convert_to_pdf=native_word_pdf
checks=[]
for src in sorted((ROOT/'交付文件').glob('*.docx')):
    folder=ROOT/'工作记录'/'逐页预览'/src.stem
    imgs=r.rasterize(str(src),str(folder),150,False,False)
    pdf=pdfplumber.open(src.with_suffix('.pdf'))
    d=Document(src)
    txt='\n'.join(p.text for p in d.paragraphs)
    extracted='\n'.join(p.extract_text() or '' for p in pdf.pages)
    norm=lambda s:re.sub(r'\s+','',s)
    expected=2 if src.name.startswith('02_') else 1
    rows=[]
    for i,p in enumerate(pdf.pages):
        spans=[s for s in p.chars if s['text'].strip()]
        rows.append({'page':i+1,'size_pt':[round(p.width,2),round(p.height,2)],'min_font_pt':min(s['size'] for s in spans),'bottom_text_pt':round(max(s['bottom'] for s in spans),2)})
    missing=[p.text for p in d.paragraphs if p.text.strip() and norm(p.text) not in norm(extracted)]
    check={'file':src.name,'pages':len(pdf.pages),'expected_pages':expected,'page_count_ok':len(pdf.pages)==expected,'A4_ok':all(abs(p.width-595.28)<1 and abs(p.height-841.89)<1 for p in pdf.pages),'missing_paragraphs':missing,'text_complete':not missing,'replacement_glyphs':'\ufffd' in extracted,'page_details':rows,'docx_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'pdf_sha256':hashlib.sha256(src.with_suffix('.pdf').read_bytes()).hexdigest(),'rendering':'Native Word PDF + canonical render_docx.py rasterize (150dpi)','visual_review':'PENDING'}
    checks.append(check)
    (folder/'PDF文字提取.txt').write_text(extracted,encoding='utf-8')
    print(json.dumps({k:check[k] for k in ['file','pages','text_complete','missing_paragraphs']},ensure_ascii=False))
(ROOT/'工作记录'/'文件检查结果.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
