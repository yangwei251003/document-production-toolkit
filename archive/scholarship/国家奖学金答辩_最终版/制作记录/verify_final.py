# Sanitized historical source; synthetic prose and portable workspace bindings.
from document_toolkit.runtime import enter_archive, workspace_root, source_path
enter_archive('scholarship/国家奖学金答辩_最终版/制作记录/verify_final.py')
from pathlib import Path
import json,hashlib,io
from PIL import Image
from pptx import Presentation
import pymupdf as fitz
root=workspace_root()
out=root/'国家奖学金答辩_最终版'
sha=lambda b:hashlib.sha256(b).hexdigest()
plan=json.loads((root/'答辩制作_16页定稿/逐页内容与素材映射.json').read_text(encoding='utf-8'))
prs=Presentation(out/'示例同学_国家奖学金答辩_16页最终版.pptx')
pdf=fitz.open(out/'示例同学_国家奖学金答辩_16页最终版.pdf')
assert len(prs.slides)==len(pdf)==16
pages=[]
for i,s in enumerate(prs.slides,1):
 p=out/'整页图片'/f'{i:02}.png';raw=p.read_bytes();im=Image.open(io.BytesIO(raw))
 assert im.size==(2560,1440)
 assert len(s.shapes)==1
 pic=s.shapes[0]
 assert sha(pic.image.blob)==sha(raw)
 assert pic.left==pic.top==0 and pic.width==prs.slide_width and pic.height==prs.slide_height
 assert s.notes_slide.notes_text_frame.text==plan['slides'][i-1]['speech']
 pp=pdf[i-1];assert tuple(pp.rect)==(0,0,960,540)
 images=pp.get_images();assert len(images)==1
 extracted=pdf.extract_image(images[0][0])
 pi=Image.open(io.BytesIO(extracted['image'])).convert('RGB')
 assert pi.size==im.size and sha(pi.tobytes())==sha(im.convert('RGB').tobytes())
 pages.append({'page':i,'png_sha256':sha(raw),'size':list(im.size),'ppt_image_match':True,'pdf_image_match':True,'speech_match':True})
manifest=json.loads((out/'制作记录/真实素材置入记录.json').read_text(encoding='utf-8'))
for a in manifest:assert sha((root/a['source']).read_bytes())==a['source_sha256']
native=list((out/'生图原稿').glob('*.png'));assert len(native)==16
report={'result':'passed','pages':pages,'source_insertions':len(manifest),'unique_sources':len({a['asset'] for a in manifest}),'original_source_hashes_match':True,'native_imagegen_pages':len(native),'files':{p.name:sha(p.read_bytes()) for p in out.iterdir() if p.suffix in ['.pptx','.pdf']}}
(out/'制作记录/最终验收.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k not in ['pages','files']},ensure_ascii=False))
