# Sanitized historical source; synthetic prose and portable workspace bindings.
from document_toolkit.runtime import enter_archive, workspace_root, source_path
enter_archive('scholarship/国家奖学金答辩_最终版/制作记录/assemble_final.py')
from pathlib import Path
import json,html,hashlib,shutil
from urllib.parse import quote
from PIL import Image,ImageDraw,ImageFont
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches
import fitz
root=workspace_root();out=root/'国家奖学金答辩_最终版';record=out/'制作记录';final=out/'整页图片'
assets={a['id']:a for a in json.loads((root/'整理说明与索引/素材索引.json').read_text(encoding='utf-8'))}
plan=json.loads((root/'答辩制作_16页定稿/逐页内容与素材映射.json').read_text(encoding='utf-8'))
# Source rectangles crop only external chrome/bars/signature; originals remain unchanged.
crops={'P023':[0,0,1122,1225],'P026':[0,885,1080,610],'P027':[42,60,1830,960],'P029':[150,636,780,1110],'P034':[25,80,1870,845],'P005':[0,100,1920,920]}
# Artwork coordinates use the generated 1672x941 reference grid. Every photo and certificate is reinserted from source.
placements={
1:[('P023',937,55,688,776,'cover')],
2:[('P025',40,201,310,403,'cover'),('P021',364,201,307,403,'cover'),('P033',683,201,308,403,'cover'),('P029',1003,201,309,431,'contain'),('P032',1326,201,308,403,'cover')],
3:[('P017',744,40,872,550,'cover'),('P025',744,604,378,234,'cover')],
4:[('P025',47,187,670,465,'cover'),('P030',737,187,464,285,'cover'),('P006',1227,155,393,500,'contain')],
5:[('P010',72,203,758,275,'cover'),('P008',843,203,759,275,'cover'),('H030',72,594,500,238,'contain'),('H038',592,594,490,238,'contain')],
6:[('P005',541,244,575,466,'contain'),('P021',1151,241,480,517,'cover')],
7:[('P005',73,261,986,523,'contain'),('H015',1108,407,497,360,'contain')],
8:[('P021',58,270,756,505,'cover'),('H008',863,447,258,354,'contain'),('H005',1151,453,463,335,'contain')],
9:[('H035',42,244,454,449,'contain'),('H048',519,242,326,453,'contain'),('H008',874,244,341,474,'contain'),('H005',1247,244,385,459,'contain')],
10:[('P033',694,37,891,436,'cover'),('P026',694,500,891,302,'cover')],
11:[('H002',80,285,507,547,'contain'),('H012',605,285,495,547,'contain'),('H006',1127,261,475,267,'contain'),('H026',1127,584,475,248,'contain')],
12:[('P029',535,102,444,631,'contain'),('P027',1018,100,496,315,'cover'),('P034',1018,467,601,314,'cover')],
13:[('P029',101,201,443,554,'contain'),('H039',689,227,495,264,'contain'),('H013',628,512,458,272,'contain'),('H037',1125,512,465,272,'contain')],
14:[('P032',700,35,868,502,'cover'),('P028',1246,357,372,454,'cover')],
15:[('P032',86,269,855,445,'cover'),('P028',978,93,607,350,'cover'),('P003',979,509,606,230,'cover')],
16:[('P023',995,45,585,506,'cover'),('P025',237,553,407,265,'cover'),('P021',673,553,405,265,'cover'),('P032',1105,553,493,265,'cover')]
}
logos={1:(76,28,444,101),5:(70,7,321,84),6:(58,22,411,92),8:(62,11,335,91),9:(44,12,345,80),10:(76,18,445,55),11:(68,12,350,90),13:(74,9,375,84),14:(65,22,418,85)}
navchap={1:'封面',2:'目录',3:'德',4:'德',5:'德',6:'智',7:'智',8:'智',9:'智',10:'体',11:'体',12:'美',13:'美',14:'劳',15:'劳',16:'结束'}
manifest=[]
def src_url(aid):return (root/assets[aid]['new']).as_uri()
def photo(aid,x,y,w,h,mode='contain',crop=None):
 a=assets[aid];fw,fh=a['final_width'],a['final_height'];box=crop or crops.get(aid,[0,0,fw,fh]);cx,cy,cw,ch=box
 clip=f'{aid}-{x}-{y}';aspect='xMidYMid slice' if mode=='cover' else 'xMidYMid meet'
 manifest.append({'asset':aid,'source':a['new'],'source_sha256':a['final_sha256'],'box':[x,y,w,h],'source_crop':box,'fit':mode})
 return f'<div class="photo" style="left:{x}px;top:{y}px;width:{w}px;height:{h}px"><svg viewBox="{cx} {cy} {cw} {ch}" preserveAspectRatio="{aspect}"><defs><clipPath id="{clip}"><rect x="{cx}" y="{cy}" width="{cw}" height="{ch}"/></clipPath></defs><image href="{src_url(aid)}" width="{fw}" height="{fh}" clip-path="url(#{clip})"/></svg></div>'
css='''*{box-sizing:border-box}body{margin:0;background:#F6F3EC}.stage{position:relative;width:1672px;height:941px;overflow:hidden;background:#F6F3EC}.base{position:absolute;width:100%;height:100%;inset:0}.photo{position:absolute;overflow:hidden;background:#F6F3EC}.photo svg{display:block;width:100%;height:100%;overflow:hidden}.brand{position:absolute;background:#F6F3EC;padding:5px 10px}.brand img{width:100%;height:100%;object-fit:contain}.footer{position:absolute;left:0;right:0;bottom:0;height:109px;background:#F6F3EC;border-top:1px solid #c9b996;font:22px 'Microsoft YaHei',sans-serif;color:#123F37;display:flex;align-items:center;padding:0 72px;gap:32px}.nav{display:flex;gap:30px}.nav span{padding:6px 10px}.nav .active{background:#123F37;color:#fff;border-radius:50%}.bar{height:3px;flex:1;background:#b9c5bb;margin:0 26px}.fill{height:100%;background:#B58B4D}.counter{font:29px Georgia,serif;white-space:nowrap}.name{font-size:18px;color:#547067}.card{position:absolute;background:#F6F3EC;color:#123F37;font-family:'Microsoft YaHei',sans-serif;padding:20px;border:1px solid #B58B4D}.mini{position:absolute;font:16px 'Microsoft YaHei',sans-serif;color:#586b60}'''
pages=[]
for n in range(1,17):
 source=next((out/'生图原稿').glob(f'{n:02}*.png'));start=len(manifest)
 body=f'<div class="stage"><img class="base" src="{source.as_uri()}">'
 for pl in placements[n]:body+=photo(*pl)
 if n in logos:
  x,y,w,h=logos[n];body+=f'<div class="brand" style="left:{x}px;top:{y}px;width:{w}px;height:{h}px"><img src="{src_url("P024")}"></div>'
 # Unrequested generated motto / campus sketch on contents removed with a clean panel.
 if n==2:body+='<div style="position:absolute;left:1230px;top:0;width:442px;height:176px;background:#F6F3EC"></div>'
 if n==5:
  body+='<div class="card" style="left:1102px;top:594px;width:500px;height:252px"><div style="font-size:32px;text-align:center">五好班集体</div><div style="font-size:20px;text-align:center;margin-top:10px">24合成示例内容：通过项目实践展示文档制作与检查流程</div><div style="font-size:15px;margin-top:10px">原名单局部</div></div>'
  body+=photo('H025',1118,741,300,33,'contain',[295,178,378,29])
  body+=photo('H025',1420,741,169,33,'contain',[1198,178,269,29])
 if n==1:
  body+='<div style="position:absolute;left:40px;top:800px;width:885px;height:32px;background:#123F37;color:#dbbe83;font:25px SimSun;letter-spacing:3px;padding:2px 42px">合成示例内容：通过项目实践展示文档制作与检查流程</div>'
 if n==4:
  body+='<div style="position:absolute;left:40px;top:802px;width:1070px;height:30px;background:#123F37;color:#dbbe83;font:23px SimSun;padding:2px 22px">合成示例内容：通过项目实践展示文档制作与检查流程</div>'
 if n==8:
  for x,w,label in [(862,263,'广东省赛 二等奖'),(1150,468,'全国总决赛 优秀奖')]:
   body+=f'<div style="position:absolute;left:{x}px;top:803px;width:{w}px;height:29px;background:#F6F3EC;color:#123F37;text-align:center;font:22px SimSun;line-height:29px">{label}</div>'
 if n==13:
  body+='<div style="position:absolute;left:95px;top:766px;width:455px;height:66px;background:#F6F3EC;color:#123F37;font:25px SimSun;line-height:31px">《拾光记忆》<br>作者：示例同学、林纳川</div>'
 # Clarify factual labels and footer independently of image-model rendering.
 nav=''.join(f'<span class="{"active" if navchap[n]==c else ""}">{c}</span>' for c in '德智体美劳')
 body+=f'<div class="footer"><span class="name">合成示例内容：通过项目实践展示文档制作与检查流程</span><div class="nav">{nav}</div><div class="bar"><div class="fill" style="width:{n / 16 * 100}%"></div></div><span class="counter">{n:02} / 16</span></div></div>'
 doc=f'<!doctype html><meta charset="utf-8"><style>{css}</style>{body}';p=record/f'page_{n:02}.html';p.write_text(doc,encoding='utf-8');pages.append(p)
 for v in manifest[start:]:v['page']=n
(record/'真实素材置入记录.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
with sync_playwright() as w:
 browser=w.chromium.launch(executable_path=source_path('msedge.exe'),headless=True)
 page=browser.new_page(viewport={'width':1672,'height':941},device_scale_factor=2560/1672)
 for n,p in enumerate(pages,1):
  page.goto(p.as_uri());page.wait_for_load_state('networkidle');page.locator('.stage').screenshot(path=str(final/f'{n:02}.png'))
 browser.close()
# Uniform exact 16:9 raster document exports; imagegen originals are 1672x941, source photos stay high resolution in compositor.
for p in final.glob('*.png'):
 with Image.open(p) as im:
  if im.size!=(2560,1440):im.resize((2560,1440),Image.Resampling.LANCZOS).save(p)
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
for n in range(1,17):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.shapes.add_picture(str(final/f'{n:02}.png'),0,0,width=prs.slide_width,height=prs.slide_height)
 s.notes_slide.notes_text_frame.text=plan['slides'][n-1]['speech']
prs.save(out/'示例同学_国家奖学金答辩_16页最终版.pptx')
pdf=fitz.open()
for n in range(1,17):
 p=pdf.new_page(width=960,height=540);p.insert_image(p.rect,filename=str(final/f'{n:02}.png'))
pdf.save(out/'示例同学_国家奖学金答辩_16页最终版.pdf',deflate=True);pdf.close()
# Contact sheet of final slides for whole-deck inspection.
sheet=Image.new('RGB',(1600,1000),'#e7e6df');d=ImageDraw.Draw(sheet);font=ImageFont.truetype(source_path('msyh.ttc'),17)
for i in range(16):
 im=Image.open(final/f'{i+1:02}.png');im.thumbnail((384,216));x=(i%4)*400+8;y=(i//4)*250+8;sheet.paste(im,(x,y));d.text((x,y+219),f'{i+1:02} / 16',font=font,fill='#123F37')
sheet.save(out/'16页总览.jpg',quality=94)
shutil.copy2(root/'答辩制作_16页定稿/16页配套答辩稿与问答提纲.md',out/'配套答辩稿.md')
print(json.dumps({'pages':16,'size':[2560,1440],'source_insertions':len(manifest),'unique_sources':len({x['asset'] for x in manifest})},ensure_ascii=False))
