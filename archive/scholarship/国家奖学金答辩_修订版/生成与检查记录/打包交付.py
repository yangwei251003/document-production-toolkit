# Sanitized historical source; synthetic prose and portable workspace bindings.
from document_toolkit.runtime import enter_archive, workspace_root, source_path
enter_archive('scholarship/国家奖学金答辩_修订版/生成与检查记录/打包交付.py')
from pathlib import Path
import json,zipfile,hashlib
from PIL import Image,ImageDraw,ImageFont
from pptx import Presentation
from pptx.util import Inches
import fitz
root=workspace_root()
out=root/'国家奖学金答辩_修订版'
nums=[1,3,4,5,6,9,13,14,15,16]
paths=[out/'整页图片'/f'{n:02}.png' for n in nums]
assert all(p.exists() for p in paths)
with zipfile.ZipFile(out/'10页修订图片.zip','w',zipfile.ZIP_DEFLATED) as z:
 for n,p in zip(nums,paths):z.write(p,f'{n:02}.png')
doc=fitz.open()
for p in paths:
 page=doc.new_page(width=960,height=540);page.insert_image(page.rect,filename=str(p),keep_proportion=False)
doc.save(out/'10页修订预览.pdf',deflate=True);doc.close()
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
for n in range(1,17):
 p=(out/'整页图片'/f'{n:02}.png') if n in nums else root/'国家奖学金答辩_最终版/整页图片'/f'{n:02}.png'
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_picture(str(p),0,0,width=prs.slide_width,height=prs.slide_height)
prs.save(out/'示例同学_国家奖学金答辩_16页修订版.pptx')
assert len(Presentation(out/'示例同学_国家奖学金答辩_16页修订版.pptx').slides)==16
sheet=Image.new('RGB',(1600,2500),'#eae7df');draw=ImageDraw.Draw(sheet);font=ImageFont.truetype(source_path('msyh.ttc'),24)
for i,(n,p) in enumerate(zip(nums,paths)):
 im=Image.open(p).convert('RGB');im.thumbnail((770,434));x=15+(i%2)*800;y=15+(i//2)*500;sheet.paste(im,(x,y));draw.text((x,y+444),f'第 {n:02} 页',font=font,fill='#123f37')
sheet.save(out/'10页修订总览.jpg',quality=94)
manifest={str(n):{'file':str(p),'size':Image.open(p).size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for n,p in zip(nums,paths)}
(out/'生成与检查记录/交付校验.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print('合成示例内容：通过项目实践展示文档制作与检查流程')
