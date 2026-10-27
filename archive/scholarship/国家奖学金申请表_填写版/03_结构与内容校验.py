# Sanitized historical source; synthetic prose and portable workspace bindings.
from document_toolkit.runtime import enter_archive, workspace_root, source_path
enter_archive('scholarship/国家奖学金申请表_填写版/03_结构与内容校验.py')
from docx import Document
from pathlib import Path
import zipfile
p=Path(__file__).resolve().parent/'2025－2026学年国家奖学金申请审批表_示例同学_24电气工程本科5班.docx'
d=Document(p)
texts="\n".join(x.text for x in d.paragraphs)+"\n"+"\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
with zipfile.ZipFile(p) as z:
    xml=z.read("word/document.xml").decode("utf-8")
app='合成示例内容：通过项目实践展示文档制作与检查流程'
rec='合成示例内容：通过项目实践展示文档制作与检查流程'
checks={
"tables":len(d.tables),
"sections":len(d.sections),
"page_breaks":xml.count('w:type="page"'),
"replacement_chars":texts.count("\ufffd"),
"application_occurrences":texts.count(app),
"recommendation_occurrences":texts.count(rec),
"application_chars":len(app),
"recommendation_chars":len(rec),
"file_bytes":p.stat().st_size
}
print(checks)
assert checks["tables"]==7
assert checks["page_breaks"]==1
assert checks["replacement_chars"]==0
assert checks["application_occurrences"]>=1
assert checks["recommendation_occurrences"]>=1
assert checks["application_chars"]<=200
assert checks["recommendation_chars"]<=100
print("PASS")

