# Sanitized session source; synthetic fixture values and explicit execution gate.
from document_toolkit.runtime import enter_archive
from document_toolkit.session_runtime import session_path, fixture_value
enter_archive('session/2026-10-07-21-33-31/.参赛信息表.ref/build.py')
# -*- coding: utf-8 -*-
try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


# 商务蓝配色
XL_PRIMARY = xl_color("#4472C4")
XL_SUBHEAD = xl_color("#D9E2F3")
XL_EMPH = xl_color("#2F5597")
XL_BORDER = xl_color("#BFBFBF")

thin = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

FONT_TITLE = Font(name="微软雅黑", size=14, bold=True, color=XL_EMPH)
FONT_SUB = Font(name="微软雅黑", size=11, bold=True, color=XL_EMPH)
FONT_HEAD = Font(name="微软雅黑", size=10, bold=True, color=xl_color("#FFFFFF"))
FONT_BODY = Font(name="微软雅黑", size=10, color=xl_color("#1A1A1A"))

FILL_HEAD = PatternFill("solid", fgColor=XL_PRIMARY)
FILL_SUB = PatternFill("solid", fgColor=XL_SUBHEAD)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

COMPETITIONS = [
    '合成示例段落：用于验证文档布局与内容转换流程',
    '合成示例段落：用于验证文档布局与内容转换流程',
]

MEMBERS = fixture_value('tables.table_1')

CLASS_NAME = "24电气工程本科5班"
MAJOR = "电气工程及其自动化"
COLLEGE = "人工智能与大数据学院"

wb = openpyxl.Workbook()
wb.properties.title = "参赛信息表"

ws = wb.active
ws.title = "参赛信息"

# ---- 锚点 ----
# 第1行：标题（合并 A1:G1）
# 第3行：区块标题「一、比赛信息」（合并 A3:G3）
# 第4行：表头  A4序号 B4比赛名称
# 第5-6行：数据
# 第8行：区块标题「二、参赛人员」（合并 A8:G8）
# 第9行：表头 A9序号 B9角色 C9姓名 D9班级 E9专业 F9学院 G9参加赛事
# 第10-14行：数据

ws.merge_cells("A1:G1")
ws["A1"] = "竞赛参赛信息表"
ws["A1"].font = FONT_TITLE
ws["A1"].alignment = CENTER
ws.row_dimensions[1].height = 30

# 一、比赛信息
ws.merge_cells("A3:G3")
ws["A3"] = "一、比赛信息"
ws["A3"].font = FONT_SUB
ws["A3"].fill = FILL_SUB
ws["A3"].alignment = LEFT

for col, name in zip("AB", ["序号", "比赛名称"]):
    ws[f"{col}4"] = name
    ws[f"{col}4"].font = FONT_HEAD
    ws[f"{col}4"].fill = FILL_HEAD
    ws[f"{col}4"].alignment = CENTER
    ws[f"{col}4"].border = BORDER

for i, name in enumerate(COMPETITIONS):
    r = 5 + i
    ws[f"A{r}"] = i + 1
    ws[f"B{r}"] = name
    for col in "AB":
        ws[f"{col}{r}"].font = FONT_BODY
        ws[f"{col}{r}"].border = BORDER
    ws[f"A{r}"].alignment = CENTER
    ws[f"B{r}"].alignment = LEFT

# 二、参赛人员
ws.merge_cells("A8:G8")
ws["A8"] = "二、参赛人员"
ws["A8"].font = FONT_SUB
ws["A8"].fill = FILL_SUB
ws["A8"].alignment = LEFT

heads = ["序号", "角色", "姓名", "班级", "专业", "学院", "参加赛事"]
for idx, name in enumerate(heads):
    col = get_column_letter(idx + 1)
    ws[f"{col}9"] = name
    ws[f"{col}9"].font = FONT_HEAD
    ws[f"{col}9"].fill = FILL_HEAD
    ws[f"{col}9"].alignment = CENTER
    ws[f"{col}9"].border = BORDER

for i, (role, name) in enumerate(MEMBERS):
    r = 10 + i
    values = [i + 1, role, name, CLASS_NAME, MAJOR, COLLEGE,
              "、".join(COMPETITIONS)]
    for idx, v in enumerate(values):
        col = get_column_letter(idx + 1)
        ws[f"{col}{r}"] = v
        ws[f"{col}{r}"].font = FONT_BODY
        ws[f"{col}{r}"].border = BORDER
        if idx in (0, 1, 2):
            ws[f"{col}{r}"].alignment = CENTER
        else:
            ws[f"{col}{r}"].alignment = LEFT

# 合计行（引用明细，人数统计）
total_row = 10 + len(MEMBERS)
ws[f"A{total_row}"] = "合计"
ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
ws[f"C{total_row}"] = f"{len(MEMBERS)} 人"
for col in "ABCDEFG":
    cell = ws[f"{col}{total_row}"]
    cell.font = Font(name="微软雅黑", size=10, bold=True, color=xl_color("#FFFFFF"))
    cell.fill = PatternFill("solid", fgColor=XL_EMPH)
    cell.border = BORDER
    cell.alignment = CENTER
ws[f"C{total_row}"].alignment = CENTER

widths = {"A": 6, "B": 8, "C": 10, "D": 22, "E": 22, "F": 22, "G": 40}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

ws.row_dimensions[3].height = 22
ws.row_dimensions[8].height = 22
ws.row_dimensions[9].height = 24
for r in range(10, total_row):
    ws.row_dimensions[r].height = 30

ws.sheet_view.showGridLines = False

OUT = session_path('input_1.xlsx')
wb.save(OUT)
print("saved:", OUT)
