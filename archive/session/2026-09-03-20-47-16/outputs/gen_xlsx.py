# Sanitized session source; synthetic fixture values and explicit execution gate.
from document_toolkit.runtime import enter_archive
from document_toolkit.session_runtime import session_path, fixture_value
enter_archive('session/2026-09-03-20-47-16/outputs/gen_xlsx.py')
# -*- coding: utf-8 -*-
fixture_value('texts.4bab3f31_1')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
PHASE_FILL = PatternFill("solid", fgColor="DDEBF7")
NOTE_FONT = Font(name="微软雅黑", size=9, italic=True, color="808080")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

def style_body(ws, r0, r1, ncols, wrap=True):
    for r in range(r0, r1 + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            if wrap:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def write_rows(ws, header, rows, start=1, widths=None, phase_col=None):
    ws.append(header)
    style_header(ws, start, len(header))
    for row in rows:
        ws.append(row)
    style_body(ws, start + 1, start + len(rows), len(header))
    # 阶段行浅色
    if phase_col:
        for i, row in enumerate(rows, start + 1):
            if row[phase_col - 1] and row[phase_col - 1] not in ("",) and "阶段" in str(row[phase_col - 1]):
                for c in range(1, len(header) + 1):
                    ws.cell(row=i, column=c).fill = PHASE_FILL
                    ws.cell(row=i, column=c).font = Font(name="微软雅黑", size=10, bold=True)
    if widths:
        set_widths(ws, widths)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)

# ============================================================ 03 计划表
wb = Workbook()
ws = wb.active
ws.title = "30天执行计划表"
header = fixture_value('tables.xlsx_header')
rows = fixture_value('tables.xlsx_rows')
write_rows(ws, header, rows, widths=[14, 8, 40, 8, 8, 12, 14, 34, 8, 10], phase_col=1)
wb.save(session_path('output_29.xlsx'))
print("03 saved")

# ============================================================ 04 多维表格
wb2 = Workbook()

# Sheet1 任务总表
ws1 = wb2.active
ws1.title = "任务总表"
write_rows(ws1, header, rows, widths=[14, 8, 40, 8, 8, 12, 14, 34, 8, 10], phase_col=1)

# Sheet2 人员分工
ws2 = wb2.create_sheet("人员分工RACI")
h2 = fixture_value('tables.xlsx_h2')
r2 = fixture_value('tables.xlsx_r2')
write_rows(ws2, h2, r2, widths=[16, 12, 20, 46, 16, 20, 22])

# Sheet3 物料清单
ws3 = wb2.create_sheet("物料清单")
h3 = fixture_value('tables.xlsx_h3')
r3 = fixture_value('tables.xlsx_r3')
write_rows(ws3, h3, r3, widths=[22, 18, 26, 10, 12, 12, 10])

# Sheet4 预算表
ws4 = wb2.create_sheet("预算表")
h4 = fixture_value('tables.xlsx_h4')
r4 = fixture_value('tables.xlsx_r4')
write_rows(ws4, h4, r4, widths=[14, 34, 14, 14, 10, 24])

# Sheet5 风险登记
ws5 = wb2.create_sheet("风险登记册")
h5 = fixture_value('tables.xlsx_h5')
r5 = fixture_value('tables.xlsx_r5')
write_rows(ws5, h5, r5, widths=[24, 8, 10, 44, 10, 22])

# Sheet6 KPI看板
ws6 = wb2.create_sheet("KPI数据周报")
h6 = fixture_value('tables.xlsx_h6')
r6 = fixture_value('tables.xlsx_r6')
write_rows(ws6, h6, r6, widths=[12, 14, 12, 12, 10, 12, 10, 12, 10, 26])

# Sheet7 积分奖品
ws7 = wb2.create_sheet("积分奖品体系")
h7 = fixture_value('tables.xlsx_h7')
r7 = fixture_value('tables.xlsx_r7')
write_rows(ws7, h7, r7, widths=[26, 12, 30, 14, 30, 10])

wb2.save(session_path('output_30.xlsx'))
print("04 saved")
