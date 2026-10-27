"""Run preserved workbook code and pure reader helpers against synthetic fixtures only."""
import argparse,ast,hashlib,html,json,re,runpy,shutil,zipfile
from pathlib import Path
from .session_runtime import fixture_value,synthetic_workbook_session

ROOT=Path(__file__).resolve().parents[1]

def reader_functions():
    """Compile ONLY the three existing pure definitions; no module IO/top-level code."""
    source=ROOT/'archive/session/2026-10-04-22-46-18/build_reader.py'
    parsed=ast.parse(source.read_text(encoding='utf-8'))
    names={'md_inline','md_to_html','human_size'}
    selected=[n for n in parsed.body if isinstance(n,ast.FunctionDef) and n.name in names]
    if len(selected)!=3:raise ValueError('Reader helper source definition mismatch')
    namespace={'re':re,'html':html,'fixture_value':fixture_value}
    exec(compile(ast.Module(body=selected,type_ignores=[]),str(source),'exec'),namespace)
    return namespace

def run(output):
    from openpyxl import load_workbook
    output=Path(output).resolve()
    if output.exists():raise ValueError('Choose a new output directory; existing files are never overwritten.')
    output.mkdir(parents=True)
    source=ROOT/'archive/session/2026-09-03-20-47-16/outputs/gen_xlsx.py'
    with synthetic_workbook_session(output):runpy.run_path(str(source),run_name='__synthetic_workbook__')
    generated=sorted(output.glob('*.xlsx'))
    if len(generated)!=2:raise ValueError('Original generator did not produce its two workbook variants')
    review=output/'review';review.mkdir();(review/'downloads').mkdir()
    workbooks=[];sections=[]
    for index,path in enumerate(generated):
        workbook=load_workbook(path,data_only=False)
        record={'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheets':[]}
        shutil.copy2(path,review/'downloads'/path.name)
        tables=[]
        for sheet in workbook:
            values=list(sheet.values)
            assert values and all(str(v).startswith('合成列') for v in values[0] if v is not None)
            assert sheet.freeze_panes is not None
            record['sheets'].append({'name':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'freeze_panes':sheet.freeze_panes,'styled_header':bool(sheet.cell(1,1).fill.fgColor.rgb)})
            rows=''.join('<tr>'+''.join(('<th>' if r==0 else '<td>')+html.escape('' if value is None else str(value))+('</th>' if r==0 else '</td>') for value in row)+'</tr>' for r,row in enumerate(values[:7]))
            tables.append('<section class="table-block"><h3>'+html.escape(sheet.title)+'</h3><p>'+str(sheet.max_row)+' 行 × '+str(sheet.max_column)+' 列 · 冻结 '+html.escape(sheet.freeze_panes)+'</p><div class="scroll"><table>'+rows+'</table></div></section>')
        workbooks.append(record)
        sections.append('<article><div class="kicker">ACTUAL XLSX / '+str(index+1)+'</div><h2>'+html.escape(path.name)+'</h2><a class="button" href="downloads/'+html.escape(path.name)+'">下载实际生成的合成工作簿</a>'+''.join(tables)+'</article>')
    helpers=reader_functions()
    markdown='# 合成阅读示例\n\n这是 **合成正文**，用于验证原Markdown转换函数。\n\n- 合成项目一\n- 合成项目二\n\n| 合成列A | 合成列B |\n| --- | --- |\n| 示例甲 | 示例乙 |\n'
    converted=helpers['md_to_html'](markdown)
    assert '<h1' in converted and '<strong>' in converted and '<li>' in converted
    inline=helpers['md_inline']('<script>example</script>')
    assert '<script>' not in inline
    (output/'reader-synthetic.md').write_text(markdown,encoding='utf-8')
    result={'synthetic_only':True,'original_generator':source.relative_to(ROOT).as_posix(),'original_reader_helpers':['md_inline','md_to_html','human_size'],'workbooks':workbooks,'reader_helper_checks':4,'original_io_executed':False,'office_executed':False,'network_requests':0,'real_signature_inputs':0}
    (output/'session-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>会话文档工具 · 实际合成结果</title><style>*{box-sizing:border-box}body{margin:0;background:#edf2f5;color:#203b4a;font-family:system-ui,"Microsoft YaHei",sans-serif}header{padding:48px 7vw;background:#173b4e;color:white}header p{max-width:850px;line-height:1.8;color:#cee0e6}h1{font-size:clamp(28px,4vw,44px)}.kicker{letter-spacing:.12em;font-size:12px;color:#6b93a3}header .kicker{color:#a7d8df}.stats{display:flex;gap:35px;margin-top:25px}.stats b{display:block;font-size:30px}.stats span{font-size:13px;color:#cee0e6}main{max-width:1200px;margin:25px auto;padding:0 24px}article{background:white;border:1px solid #dbe5eb;padding:30px;border-radius:12px;margin:22px 0}.button{display:inline-block;padding:10px 15px;background:#e3f0f4;border-radius:6px;color:#225e75;text-decoration:none}.table-block{margin:26px 0}.scroll{overflow-x:auto}table{border-collapse:collapse;font-size:12px;min-width:650px;width:100%}td,th{padding:9px;border:1px solid #dce6ec;text-align:left;white-space:nowrap}th{background:#244c62;color:white}p{line-height:1.7}.reader{border-left:4px solid #79aabb;padding:15px 24px;background:#f5f9fb}.reader table{min-width:0}@media(max-width:650px){header{padding:30px 24px}article{padding:22px}.stats{gap:20px}main{padding:0 14px}}</style><header><div class="kicker">DOCUMENT PRODUCTION TOOLKIT / 0.2</div><h1>保全程序，让会话成果可以复现</h1><p>原工作簿生成程序实际运行，表格内容全部改为分离的合成fixture；阅读区调用原有Markdown转换函数。未读取原文章、财务表、签名图片或私有平台运行库。</p><div class="stats"><div><b>55</b><span>累计保全程序</span></div><div><b>20</b><span>本次新增源码</span></div><div><b>2</b><span>实际生成XLSX</span></div><div><b>8</b><span>工作表</span></div></div></header><main><article><div class="kicker">ORIGINAL PURE READER FUNCTIONS</div><h2>合成Markdown → HTML</h2><div class="reader">'''+converted+'</div><p>仅调用三个原纯函数；完整资料阅读器的原数据加载与生成流程未执行。</p></article>'+''.join(sections)+'</main></html>'
    (review/'index.html').write_text(page,encoding='utf-8')
    print(json.dumps({'output':str(output),'workbooks':len(workbooks),'sheets':sum(len(x['sheets']) for x in workbooks),'reader_checks':4,'office':False,'network_requests':0},ensure_ascii=False))
    return result

def main():
    p=argparse.ArgumentParser();p.add_argument('--output',required=True);args=p.parse_args();run(args.output)

if __name__=='__main__':
    try:main()
    except (ValueError,FileNotFoundError) as error:raise SystemExit(str(error))
